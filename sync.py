"""
同步兼職資料 → jobs-data.js

規則：
1. 以原網站 /api/jobs 為底
2. 名稱帶【需要確認】的任務，改用試算表的資料
3. 班克爾區只在試算表有，直接從試算表加入（備註欄＝製作/加工需要的原料）
4. 試算表任務名稱是黃底的 → 標大推；要求寫「製作N次」「加工N次」→ 標製作／加工

用法：python sync.py
"""
import io
import json
import re
import sys
import urllib.request
from datetime import datetime
from pathlib import Path

import openpyxl

API_URL = "https://mabinogi-mobile-jobs.vtuberparrot2021.chatgpt.site/api/jobs"
SHEET_URL = "https://docs.google.com/spreadsheets/d/1J2bhEzdvVkuyCFSnW31yTX7sE30mD4oEOHE7DnrXQjc/export?format=xlsx"
OUT = Path(__file__).with_name("jobs-data.js")

UNCONFIRMED = "【需要確認】"
SHEET_REGIONS = ["堤爾克那", "杜巴頓", "庫漢", "班克爾"]
SHEET_ONLY_REGIONS = ["班克爾"]
STAR_FILL = "FFFFE599"  # 試算表「高價值兼職」的黃底

# 同一種東西不同寫法 → 統一名稱，總數才會合併計算
ALIASES = {
    "水": "裝水的瓶子",
    "一瓶裝滿水的瓶子": "裝水的瓶子",
}

# 不能共用保管箱的材料：只寫在備註，不列入總數
NOT_COUNTED = {"煉金術碎屑"}

# 試算表沒寫製作/加工，但實際需要的（使用者確認過）
CRAFT_OVERRIDES = {
    "滿懷溫暖的心意": {"type": "製作", "times": 1},
}


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read()


def plain_name(name):
    """去掉【跑腿】【大推】等前綴與標點，用來比對名稱"""
    name = re.sub(r"【[^】]*】", "", name)
    return re.sub(r"[？?！!\s]", "", name)


def alias(name):
    return ALIASES.get(name, name)


def parse_requirement(req):
    """試算表的「要求」欄 → (交付物名稱, 是否跑腿, 製作/加工, 備註)"""
    req = str(req or "").strip()
    if req.startswith("跑腿"):
        return req, True, None, ""
    craft, note = None, ""
    if "+交付" in req:
        head, req = req.split("+交付", 1)
        m = re.match(r"^(製作|加工)(\d+)次$", head)
        if m:
            craft = {"type": m.group(1), "times": int(m.group(2))}
        else:
            note = head
    m = re.match(r"^(.*?)\((.*)\)$", req)
    if m:
        req, extra = m.group(1), m.group(2)
        note = (note + " " + extra).strip()
    return req, False, craft, note


def parse_materials(text):
    """備註「4蘋果 2藥草 1 水 2糖」→ [{name, quantity}]"""
    return [
        {"name": alias(name), "quantity": int(n)}
        for n, name in re.findall(r"(\d+)\s*([^\d\s]+)", str(text or ""))
    ]


def read_sheet(raw):
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    rows = []
    for region in SHEET_REGIONS:
        if region not in wb.sheetnames:
            continue
        ws = wb[region]
        area = region
        for cells in ws.iter_rows(min_row=1, max_col=7):
            a, shop, npc, name, req, qty, memo = ([c.value for c in cells] + [None] * 7)[:7]
            if a and str(a).strip():
                area = str(a).strip()
            if not name or name == "任務名稱":
                continue
            fill = cells[3].fill
            star = bool(fill.fill_type) and fill.fgColor.rgb == STAR_FILL
            deliver, errand, craft, note = parse_requirement(req)
            name = str(name).strip()
            craft = CRAFT_OVERRIDES.get(name, craft)
            rows.append({
                "region": region, "area": area,
                "shop": str(shop or "").strip(), "npc": str(npc or "").strip(),
                "name": name, "deliver": alias(deliver),
                "qty": int(qty) if isinstance(qty, (int, float)) else None,
                "errand": errand, "craft": craft, "note": note, "star": star,
                "materials": parse_materials(memo),
            })
    return rows


def sheet_row_to_job(row):
    name = row["name"]
    if row["errand"] and not name.startswith("【"):
        name = "【跑腿】" + name
    item = {"name": row["deliver"], "quantity": 1 if row["errand"] else row["qty"]}
    counted = [m for m in row["materials"] if m["name"] not in NOT_COUNTED]
    personal = [m for m in row["materials"] if m["name"] in NOT_COUNTED]
    job = {
        "city": row["region"], "area": row["area"],
        "shop": row["shop"], "npc": row["npc"], "name": name,
        "deliverable": item,
        "materials": counted or [dict(item)],
        "afterAccept": bool(row["craft"]), "source": "sheet",
    }
    notes = []
    if row["note"]:
        notes.append(row["note"])
    if row["craft"]:
        job["craft"] = row["craft"]
        if not row["materials"]:
            notes.append("試算表沒寫原料，先以成品計算")
    if personal:
        items = "、".join(f'{m["name"]}×{m["quantity"]}' for m in personal)
        notes.append(f"另需 {items}（不能共用保管箱，各角色自備，不列入總數）")
    if notes:
        job["note"] = "；".join(notes)
    if row["star"]:
        job["star"] = True
    return job


def main():
    site = json.loads(fetch(API_URL))
    sheet = read_sheet(fetch(SHEET_URL))
    by_name = {plain_name(r["name"]): r for r in sheet}

    jobs, patched, missing, starred = [], [], [], []
    for job in site["jobs"]:
        job = dict(job)
        job.pop("id", None)
        row = by_name.get(plain_name(job["name"]))
        if UNCONFIRMED in job["name"]:
            if row:
                fixed = sheet_row_to_job(row)
                fixed["city"], fixed["area"] = job["city"], job["city"]
                patched.append(f'{job["name"]} → {row["deliver"]}×{row["qty"]}')
                job = fixed
            else:
                missing.append(job["name"])
        elif row:
            # 原網站的資料為主，只從試算表補上 大推 與 製作/加工 標記
            if row["craft"]:
                job["craft"] = row["craft"]
            if row["star"] and "【大推】" not in job["name"]:
                job["star"] = True
                starred.append(f'{job["city"]} {plain_name(job["name"])}')
        for m in job["materials"]:
            m["name"] = alias(m["name"])
        job["deliverable"]["name"] = alias(job["deliverable"]["name"])
        job.setdefault("area", job["city"])
        jobs.append(job)

    for row in sheet:
        if row["region"] in SHEET_ONLY_REGIONS:
            jobs.append(sheet_row_to_job(row))

    # id 用名稱組成，重新同步後已勾選的任務才不會跑掉
    seen = {}
    for job in jobs:
        base = f'{job["city"]}|{job["npc"]}|{plain_name(job["name"])}'
        seen[base] = seen.get(base, 0) + 1
        job["id"] = base if seen[base] == 1 else f"{base}#{seen[base]}"

    data = {
        "jobs": jobs,
        "siteUpdatedAt": site.get("updatedAt"),
        "syncedAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    OUT.write_text(
        "// 由 sync.py 產生，請勿手動修改\nwindow.JOBS_DATA = "
        + json.dumps(data, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )
    print(f"共 {len(jobs)} 筆任務 → {OUT.name}")
    for p in patched:
        print("  以試算表修正：", p)
    for s in starred:
        print("  依黃底標大推：", s)
    for m in missing:
        print("  ⚠ 試算表找不到：", m)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
